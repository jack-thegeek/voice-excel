"""Excel 成绩表读写。目标文件固定为本目录下 成绩表.xlsx。"""
from __future__ import annotations

import os
import shutil
from dataclasses import dataclass
from typing import Optional

import openpyxl

EXCEL_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "成绩表.xlsx")
SHEET_NAME = "班级成绩登记表"

# 表头行
HEADER_ROW = 2
# 数据行范围（含）
DATA_START_ROW = 3
DATA_END_ROW = 50
# 列
SEQ_COL = 1          # A 序号
NAME_COL = 2         # B 姓名
# 可修改的分数列映射（表头文本 -> 列号）
SCORE_COLS: dict[str, int] = {
    "第一单元": 3,
    "第二单元": 4,
    "第三单元": 5,
    "第四单元": 6,
    "第五单元": 7,
    "第六单元": 8,
    "期中考试": 9,
    "期末考试": 10,
}
DEFAULT_COL_NAME = "第一单元"


class ExcelError(Exception):
    pass


@dataclass
class Student:
    row: int
    seq: int
    name: str
    scores: dict[str, object]  # 列名 -> 值


def _load() -> tuple[openpyxl.Workbook, openpyxl.worksheet.worksheet.Worksheet]:
    if not os.path.exists(EXCEL_PATH):
        raise ExcelError(f"找不到成绩表: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    if SHEET_NAME not in wb.sheetnames:
        raise ExcelError(f"工作表 {SHEET_NAME!r} 不存在，现有: {wb.sheetnames}")
    return wb, wb[SHEET_NAME]


def list_students() -> list[Student]:
    """读取所有学生（按序号列非空、非统计行）。"""
    _wb, ws = _load()
    out: list[Student] = []
    for r in range(DATA_START_ROW, DATA_END_ROW + 1):
        seq = ws.cell(r, SEQ_COL).value
        if seq is None or not isinstance(seq, (int, float)):
            # 跳过统计行（"平均分"等字符串）与空行
            continue
        name = ws.cell(r, NAME_COL).value or ""
        scores = {col_name: ws.cell(r, idx).value for col_name, idx in SCORE_COLS.items()}
        out.append(Student(row=r, seq=int(seq), name=str(name), scores=scores))
    return out


def find_student(seq: int) -> Optional[Student]:
    for s in list_students():
        if s.seq == seq:
            return s
    return None


def update_score(seq: int, col_name: str, new_value) -> Student:
    """修改指定序号、指定列的分数并保存。返回更新后的学生信息。"""
    if col_name not in SCORE_COLS:
        raise ExcelError(f"未知分数列: {col_name}，可选: {list(SCORE_COLS)}")
    wb, ws = _load()
    target_row: Optional[int] = None
    for r in range(DATA_START_ROW, DATA_END_ROW + 1):
        v = ws.cell(r, SEQ_COL).value
        if isinstance(v, (int, float)) and int(v) == int(seq):
            target_row = r
            break
    if target_row is None:
        raise ExcelError(f"找不到序号 {seq} 的学生")

    col_idx = SCORE_COLS[col_name]
    ws.cell(target_row, col_idx, new_value)
    # 保持单元格格式为"常规"(General)
    ws.cell(target_row, col_idx).number_format = "General"

    # 备份原文件一次（首次写入时）
    bak = EXCEL_PATH + ".bak"
    if not os.path.exists(bak):
        shutil.copy2(EXCEL_PATH, bak)
    wb.save(EXCEL_PATH)

    # 返回最新数据
    return find_student(seq)  # type: ignore[return-value]
